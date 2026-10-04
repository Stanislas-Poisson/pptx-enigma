#!/usr/bin/env python3
"""Builds examples/bakery.pptx: a fictional presentation whose speaker notes hold voice-overs.

Everything in it is invented. It shows what the extractor reads: several speakers, bold, italic and
underline, a line break, a numbered list, a bulleted list, a link, accents, and two voice-overs in the
same notes.

    pip install python-pptx
    python3 examples/build_bakery.py
"""

import copy
import os

from pptx import Presentation
from pptx.oxml.ns import qn

SIGN = "¤"
OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bakery.pptx")

SLIDE_TITLES = [
    "Moonlight Bakery",
    "Welcome",
    "The dough",
    "The oven",
    "A customer",
    "Goodbye",
]


def marker(speaker, reference):
    return f"{SIGN} VOICE OVER ({speaker}) {reference}"


END = f"{SIGN} VOICE OVER END"

# A paragraph is a text, or a dict: runs (a list of texts or of (text, style) pairs), bullet ("ul" or "ol")
# and level.
NOTES = [
    # The notes of the first slide only give the sign.
    [SIGN],
    [
        marker("Narrator", "w01_intro"),
        {
            "runs": [
                "Welcome to the ",
                ("Moonlight Bakery", "b"),
                ". Tonight, we follow ",
                ("one", "i"),
                " loaf of bread, from the flour to the shelf.",
            ]
        },
        "It takes a little patience, and a lot of love.",
        END,
    ],
    [
        marker("Baker", "w02_dough"),
        "First, the dough. You need four things:",
        {"runs": ["500 g of flour"], "bullet": "ol"},
        {"runs": ["350 ml of water"], "bullet": "ol"},
        {"runs": ["10 g of salt"], "bullet": "ol"},
        {"runs": ["a pinch of yeast"], "bullet": "ol"},
        {"runs": ["Mix them, then ", ("wait", "bu"), ". The dough rises on its own."]},
        marker("Narrator", "w03_wait"),
        {"runs": ["Waiting is the hardest part.", "\n", "Two hours, without touching it."]},
        END,
    ],
    [
        marker("Baker", "w04_oven"),
        "The oven must be very hot. Remember:",
        {"runs": ["240 degrees"], "bullet": "ul"},
        {"runs": ["a tray of water at the bottom"], "bullet": "ul"},
        {"runs": ["never open the door in the first ten minutes"], "bullet": "ul", "level": 1},
        {"runs": ["Bake for ", ("thirty", "b"), " minutes."]},
        END,
    ],
    [
        marker("Customer", "w05_order"),
        {
            "runs": [
                "Bonjour ! Je voudrais une ",
                ("baguette", "b"),
                ", s'il vous plaît, et un croissant à l'été.",
            ]
        },
        {"runs": ["The menu is on ", ("our website", "link:https://example.com/menu"), "."]},
        END,
    ],
    [
        marker("Narrator", "w06_end"),
        "And that is how a loaf of bread is made.",
        {"runs": [("Thank you for watching.", "bi")]},
        END,
    ],
]


def add_runs(paragraph, runs):
    for run in runs:
        if isinstance(run, str) and run == "\n":
            paragraph.add_line_break()
            continue

        text, style = (run, "") if isinstance(run, str) else run
        part = paragraph.add_run()
        part.text = text

        if style.startswith("link:"):
            part.hyperlink.address = style[len("link:"):]
            continue

        part.font.bold = "b" in style or None
        part.font.italic = "i" in style or None
        part.font.underline = "u" in style or None


def set_bullet(paragraph, kind, level):
    properties = paragraph._p.get_or_add_pPr()
    properties.set("lvl", str(level))
    properties.set("marL", str(342900 + 342900 * level))
    properties.set("indent", "-342900")

    for child in list(properties):
        properties.remove(child)

    font = properties.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    properties.append(font)
    bullet = (
        properties.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
        if kind == "ol"
        else properties.makeelement(qn("a:buChar"), {"char": "•"})
    )
    properties.append(bullet)


def write_notes(slide, lines):
    frame = slide.notes_slide.notes_text_frame
    first = True

    for line in lines:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False

        if isinstance(line, str):
            add_runs(paragraph, [line])
            continue

        add_runs(paragraph, line["runs"])

        if "bullet" in line:
            set_bullet(paragraph, line["bullet"], line.get("level", 0))


def main():
    presentation = Presentation()
    layout = presentation.slide_layouts[5]

    for title, lines in zip(SLIDE_TITLES, NOTES):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        write_notes(slide, lines)

    presentation.save(OUTPUT)
    print(f"written {OUTPUT}")


if __name__ == "__main__":
    main()
